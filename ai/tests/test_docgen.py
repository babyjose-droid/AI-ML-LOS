import io

import docx
from fastapi.testclient import TestClient

from app.docgen import docx_to_html, sanitise
from app.main import app

client = TestClient(app)


def test_pdf_in_english_and_hindi():
    r = client.post("/v1/docs/pdf", json={"html": "<h1>Loan agreement</h1><p>ऋण समझौता</p><table><tr><td>1</td></tr></table>", "title": "Agreement"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    assert client.post("/v1/docs/pdf", json={"html": " "}).status_code == 422


def test_no_scripts_or_external_resources():
    s = sanitise('<p>x</p><script>alert(1)</script><img src="http://evil.example/a.png"><img src="data:image/png;base64,AA"><link rel="stylesheet" href="http://x">')
    assert "script" not in s and "evil" not in s and "link" not in s
    assert "data:image/png" in s


def test_word_template_keeps_fields_split_across_runs():
    d = docx.Document()
    d.add_heading("Loan Agreement", 1)
    p = d.add_paragraph("Between ")
    p.add_run("{{applicant").bold = True
    p.add_run(".name}} and the lender for {{loan.amount | rupees}}.")
    b = io.BytesIO()
    d.save(b)
    out = docx_to_html(b.getvalue())
    assert "<h1>Loan Agreement</h1>" in out["html"]
    assert "{{applicant.name}}" in out["html"]
    assert out["fields"] == ["applicant.name", "loan.amount | rupees"]
    r = client.post("/v1/docs/docx-to-html", files={"file": ("t.docx", b.getvalue(), "application/octet-stream")})
    assert r.status_code == 200 and "applicant.name" in r.json()["fields"]
    assert client.post("/v1/docs/docx-to-html", files={"file": ("t.docx", b"not a zip", "application/octet-stream")}).status_code == 415
